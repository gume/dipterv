#!/usr/bin/env python
# -*- coding: utf-8 -*-

import datetime
import glob
import os
import re
import sqlite3
import sys

from openpyxl import load_workbook


DATABASE_PATH = "/db/descriptions.sqlite"
PORTAL_STATUS_MAP = {
    "Feltöltve, tanszékvezetői jóváhagyásra vár": "undecided",
    "Konzulensi feltöltésre vár, mert a tanszékvezető elutasította": "rejected",
    "Tanszékvezető jóváhagyta": "accepted",
    "Nincs beadva, mert a tanszékvezető nem hagyta jóvá": "rejected",
}
PENDING_STATUSES = set(PORTAL_STATUS_MAP)
IMPORT_COLUMNS = (
    "neptun_code",
    "row_id",
    "filename",
    "version",
    "student_name",
    "supervisor_name",
    "supervisor_code",
    "degree",
    "faculty",
    "subject",
    "title_hu",
    "title_en",
    "portal_status",
    "portal_status_date",
)
# Initial entries of the reviewers table: (id, display name).
HUMAN_REVIEWERS = (
    ("kovacsg", "KovacsG"),
    ("guszti", "Guszti"),
    ("gume", "Gume"),
    ("pali", "Pali"),
)
EXTRACT_COLUMNS = {
    "template_status": "INTEGER",
    "page_count": "INTEGER",
    "detected_template": "TEXT",
    "header_dept_head": "TEXT",
    "thesis_type": "TEXT",
    "student_suffix": "TEXT",
    "task_intro": "TEXT",
    "supervisor_label": "TEXT",
    "dept_head_title": "TEXT",
    "footer_university": "TEXT",
    "footer_faculty": "TEXT",
    "footer_department": "TEXT",
    "missing_boilerplate_rules_json": "TEXT",
    "layout_discrepancies_json": "TEXT",
    "student_name": "TEXT",
    "student_status": "TEXT",
    "thesis_title": "TEXT",
    "problem_background": "TEXT",
    "project_goal": "TEXT",
    "tasks_json": "TEXT",
    "supervisor": "TEXT",
    "cosupervisor": "TEXT",
    "department_head": "TEXT",
    "issue_date": "TEXT",
    "university_logo_present": "INTEGER",
    "university_logo_box_2d_json": "TEXT",
    "department_logo_present": "INTEGER",
    "department_logo_box_2d_json": "TEXT",
    "signature_space_present": "INTEGER",
    "signature_space_box_2d_json": "TEXT",
    "keyword1": "TEXT",
    "keyword2": "TEXT",
    "keyword3": "TEXT",
    "keyword4": "TEXT",
    "keyword5": "TEXT",
}


def normalize_text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def parse_date(value):
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, datetime.date):
        return datetime.datetime.combine(value, datetime.time.min)
    if isinstance(value, str):
        v = value.strip()
        if not v:
            return None
        try:
            dt = datetime.datetime.fromisoformat(v)
            return dt.replace(tzinfo=None) if dt.tzinfo else dt
        except ValueError:
            return None
    return None


def normalize_date(value):
    dt = parse_date(value)
    return dt.strftime("%Y-%m-%d %H:%M:%S") if dt else None


def is_older_than_months(value, months=3, now=None):
    dt = parse_date(value)
    if dt is None:
        return False
    if now is None:
        now = datetime.datetime.now()
    cutoff = now - datetime.timedelta(days=30 * months)
    return dt < cutoff


def parse_version(filename):
    match = re.search(r"-(\d+)\.pdf$", filename or "", re.IGNORECASE)
    return int(match.group(1)) if match else None


def load_import_rows(workbook_path, data_dir="/data"):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [normalize_text(cell) for cell in next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))]
    required_headers = (
        "Hallgató Neptun kódja",
        "Hallgató neve",
        "Konzulens neve",
        "Konzulens Neptun kódja",
        "Képzés",
        "Szak",
        "Tantárgy",
        "Téma magyar címe",
        "Téma angol címe",
        "Feladatkiírás verziószáma",
        "Feladatkiírás státusza",
        "Feladatkiírás jóváhagyás/elutasítás dátuma",
        "Téma oktatói szerkesztő oldalának címe (URL) a portálon",
    )
    missing_headers = [header for header in required_headers if header not in headers]
    if missing_headers:
        workbook.close()
        raise ValueError("Workbook is missing required columns: " + ", ".join(missing_headers))
    header_indexes = {header: headers.index(header) for header in required_headers}

    records = []
    for row_id, cells in enumerate(sheet.iter_rows(min_row=5, values_only=True)):
        row = {header: cells[index] if index < len(cells) else None for header, index in header_indexes.items()}
        raw_status = normalize_text(row["Feladatkiírás státusza"])
        if raw_status not in PORTAL_STATUS_MAP:
            continue

        raw_date = row.get("Feladatkiírás jóváhagyás/elutasítás dátuma")
        if is_older_than_months(raw_date, months=3):
            continue

        neptun_code = normalize_text(row["Hallgató Neptun kódja"])
        url = normalize_text(row["Téma oktatói szerkesztő oldalának címe (URL) a portálon"])
        slug_match = re.search(r"Theses/([^/]+)/", url or "")
        if not neptun_code or not slug_match:
            print(f"WARNING: Skipping workbook row {row_id}: missing Neptun code or task URL.")
            continue

        try:
            sheet_version = int(float(row["Feladatkiírás verziószáma"]))
        except (TypeError, ValueError):
            sheet_version = None

        slug = slug_match.group(1)
        existing_files = glob.glob(os.path.join(data_dir, f"{slug}-Feladatkiiras-*.pdf"))
        if sheet_version:
            filename = f"{slug}-Feladatkiiras-{sheet_version}.pdf"
        elif existing_files:
            filename = max(existing_files, key=lambda path: parse_version(os.path.basename(path)) or 0)
            filename = os.path.basename(filename)
        else:
            print(f"WARNING: Skipping workbook row {row_id}: cannot determine PDF version.")
            continue

        version = parse_version(filename)
        if version is None:
            print(f"WARNING: Skipping workbook row {row_id}: cannot parse version from {filename}.")
            continue

        records.append({
            "neptun_code": neptun_code,
            "row_id": row_id,
            "filename": filename,
            "version": version,
            "student_name": normalize_text(row["Hallgató neve"]),
            "supervisor_name": normalize_text(row["Konzulens neve"]),
            "supervisor_code": (normalize_text(row["Konzulens Neptun kódja"]) or "").upper() or None,
            "degree": normalize_text(row["Képzés"]),
            "faculty": normalize_text(row["Szak"]),
            "subject": normalize_text(row["Tantárgy"]),
            "title_hu": normalize_text(row["Téma magyar címe"]),
            "title_en": normalize_text(row["Téma angol címe"]),
            "portal_status": PORTAL_STATUS_MAP.get(raw_status),
            "portal_status_date": normalize_date(raw_date),
            "download_url": url,
            "should_download": raw_status
            == "Feltöltve, tanszékvezetői jóváhagyásra vár",
        })

    workbook.close()
    return records


def ensure_schema(connection):
    connection.row_factory = sqlite3.Row
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS "import" (
            neptun_code TEXT NOT NULL,
            row_id INTEGER,
            filename TEXT NOT NULL UNIQUE,
            version INTEGER NOT NULL,
            student_name TEXT,
            supervisor_name TEXT,
            supervisor_code TEXT,
            degree TEXT,
            faculty TEXT,
            subject TEXT,
            title_hu TEXT,
            title_en TEXT,
            portal_status TEXT,
            portal_status_date TEXT,
            UNIQUE (neptun_code, version)
        );
        CREATE TABLE IF NOT EXISTS "reviews" (
            neptun_code TEXT NOT NULL,
            version INTEGER NOT NULL,
            reviewer TEXT NOT NULL,
            aspect TEXT NOT NULL,
            result TEXT,
            score INTEGER,
            updated_at TEXT NOT NULL DEFAULT (datetime('now')),
            PRIMARY KEY (neptun_code, version, reviewer, aspect)
        );
        CREATE TABLE IF NOT EXISTS "review_summaries" (
            neptun_code TEXT NOT NULL,
            version INTEGER NOT NULL,
            reviewer TEXT NOT NULL,
            average_score REAL NOT NULL,
            summary TEXT NOT NULL,
            updated_at TEXT NOT NULL DEFAULT (datetime('now')),
            PRIMARY KEY (neptun_code, version, reviewer)
        );
        CREATE TABLE IF NOT EXISTS "review_verdict" (
            neptun_code TEXT NOT NULL,
            version INTEGER NOT NULL,
            status TEXT,
            summary TEXT,
            updated_at TEXT NOT NULL DEFAULT (datetime('now')),
            PRIMARY KEY (neptun_code, version)
        );
        CREATE TABLE IF NOT EXISTS "reviewers" (
            id TEXT PRIMARY KEY NOT NULL,
            name TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS "supervisors" (
            neptun_code TEXT PRIMARY KEY NOT NULL,
            name TEXT NOT NULL,
            alt_name TEXT
        );
        CREATE TABLE IF NOT EXISTS "settings" (
            key TEXT PRIMARY KEY NOT NULL,
            value TEXT
        );
        CREATE TABLE IF NOT EXISTS "description_topics" (
            rank INTEGER PRIMARY KEY NOT NULL,
            topic_name TEXT NOT NULL,
            topic_description TEXT,
            thesis_count INTEGER NOT NULL,
            percentage REAL NOT NULL,
            top_keywords_json TEXT,
            updated_at TEXT NOT NULL DEFAULT (datetime('now'))
        );
        CREATE TABLE IF NOT EXISTS "description_metric_distributions" (
            metric_key TEXT PRIMARY KEY NOT NULL,
            metric_label TEXT NOT NULL,
            source_type TEXT NOT NULL,
            average_value REAL NOT NULL,
            min_value REAL,
            max_value REAL,
            total_samples INTEGER NOT NULL,
            histogram_json TEXT NOT NULL,
            updated_at TEXT NOT NULL DEFAULT (datetime('now'))
        );
        CREATE TABLE IF NOT EXISTS "description_fun_facts" (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            category TEXT NOT NULL,
            title TEXT NOT NULL,
            fact_text TEXT NOT NULL,
            highlight_stat TEXT,
            updated_at TEXT NOT NULL DEFAULT (datetime('now'))
        );
        """
    )
    connection.executemany(
        'INSERT OR IGNORE INTO "reviewers" (id, name) VALUES (?, ?)', HUMAN_REVIEWERS
    )
    ensure_extract_schema(connection)


def load_supervisor_rows(workbook_path):
    """Distinct (neptun_code, name) pairs of the supervisors listed in the workbook."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [normalize_text(cell) for cell in next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))]
    try:
        neptun_index = headers.index("Konzulens Neptun kódja")
        name_index = headers.index("Konzulens neve")
    except ValueError:
        workbook.close()
        raise ValueError("Workbook is missing the supervisor Neptun code or name column.")
    supervisors = {}
    for cells in sheet.iter_rows(min_row=5, values_only=True):
        neptun = normalize_text(cells[neptun_index]) if neptun_index < len(cells) else None
        name = normalize_text(cells[name_index]) if name_index < len(cells) else None
        if neptun and name:
            supervisors[neptun.upper()] = name
    workbook.close()
    return supervisors


def sync_supervisors(supervisors, database_path=DATABASE_PATH):
    """Add new supervisors and refresh names; manually set alt_name values are kept."""
    connection = sqlite3.connect(database_path)
    try:
        ensure_schema(connection)
        for neptun, name in supervisors.items():
            current = connection.execute(
                'SELECT name FROM supervisors WHERE neptun_code = ?', (neptun,)
            ).fetchone()
            if current is None:
                connection.execute(
                    'INSERT INTO supervisors (neptun_code, name) VALUES (?, ?)', (neptun, name)
                )
                print(f"Added supervisor {neptun}: {name}")
            elif current[0] != name:
                connection.execute(
                    'UPDATE supervisors SET name = ? WHERE neptun_code = ?', (name, neptun)
                )
                print(f"WARNING: Supervisor {neptun} renamed: {current[0]} -> {name}")
        connection.commit()
    finally:
        connection.close()


def ensure_extract_schema(connection):
    column_definitions = ",\n            ".join(
        f'"{column}" {column_type}'
        for column, column_type in EXTRACT_COLUMNS.items()
    )
    connection.execute(
        f'CREATE TABLE IF NOT EXISTS "extract" ('
        f'filename TEXT PRIMARY KEY NOT NULL, {column_definitions})'
    )
    existing = {row[1] for row in connection.execute('PRAGMA table_info("extract")')}
    for column, column_type in EXTRACT_COLUMNS.items():
        if column not in existing:
            connection.execute(f'ALTER TABLE "extract" ADD COLUMN "{column}" {column_type}')


def sync_import_records(records, database_path=DATABASE_PATH):
    connection = sqlite3.connect(database_path)
    try:
        ensure_schema(connection)
        connection.commit()
        connection.execute("BEGIN")
        for record in records:
            values = tuple(record.get(column) for column in IMPORT_COLUMNS)
            current = connection.execute(
                'SELECT * FROM "import" WHERE (neptun_code = ? AND version = ?) OR filename = ? '
                'ORDER BY CASE WHEN neptun_code = ? AND version = ? THEN 0 ELSE 1 END LIMIT 1',
                (record["neptun_code"], record["version"], record["filename"], record["neptun_code"], record["version"]),
            ).fetchone()

            if current is None:
                placeholders = ", ".join("?" for _ in IMPORT_COLUMNS)
                columns = ", ".join(IMPORT_COLUMNS)
                connection.execute(
                    f'INSERT INTO "import" ({columns}) VALUES ({placeholders})', values
                )
                print(
                    f"Added import record: row {record['row_id']}, "
                    f"version {record['version']}, {record['filename']}"
                )
            else:
                missing = []
                changed = []
                for column, new_value in zip(IMPORT_COLUMNS, values):
                    old_value = current[column]
                    if old_value != new_value:
                        if old_value in (None, "") and new_value not in (None, ""):
                            missing.append(column)
                        else:
                            changed.append(column)
                if missing:
                    print(
                        f"Added missing import data for {record['neptun_code']} "
                        f"version {record['version']}: {', '.join(missing)}"
                    )
                if changed:
                    print(
                        f"WARNING: Import data differs for {record['neptun_code']} "
                        f"version {record['version']}; overwriting: {', '.join(changed)}"
                    )
                assignments = ", ".join(f"{column} = ?" for column in IMPORT_COLUMNS)
                connection.execute(
                    f'UPDATE "import" SET {assignments} WHERE neptun_code = ? AND version = ?',
                    values + (current["neptun_code"], current["version"]),
                )
                if (current["neptun_code"], current["version"]) != (record["neptun_code"], record["version"]):
                    for table in ("reviews", "review_summaries", "review_verdict"):
                        connection.execute(
                            f"UPDATE {table} SET neptun_code = ?, version = ? WHERE neptun_code = ? AND version = ?",
                            (record["neptun_code"], record["version"], current["neptun_code"], current["version"]),
                        )

        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Usage: create_db.py <workbook.xlsx>")
    records = load_import_rows(sys.argv[1])
    sync_import_records(records)
    sync_supervisors(load_supervisor_rows(sys.argv[1]))

if __name__ == "__main__":
    main()